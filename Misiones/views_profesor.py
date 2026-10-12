"""Vistas del panel del profesor."""
import json
import random
from datetime import timedelta

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Q, Sum, Value
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import gamificacion as juego
from . import minijuegos
from .forms import (
    AltaMasivaForm, AlumnoForm, ArticuloTiendaForm, ClaseForm, LogroForm, MisionDiariaForm, MisionForm,
    MundoForm, NivelForm, OtorgarLogroForm, separar_nombre,
)
from .models import (
    Alumno, ArticuloTienda, Clase, CompraAlumno, Comportamiento, Equipo, Juego, Logro, LogroAlumno, MetaGrupo, Mision,
    MisionDiaria, Movimiento, Mundo, Nivel, OpcionRespuesta, Pregunta, ProgresoMision,
)
from .templatetags.nextlevel import COLORES_CLASE, avatar_url


def _siguiente(request, por_defecto):
    destino = request.POST.get('next') or request.GET.get('next')
    if destino and url_has_allowed_host_and_scheme(destino, allowed_hosts={request.get_host()}):
        return destino
    return por_defecto


def _es_modal(request):
    """Las ventanas emergentes del panel piden los formularios con la cabecera X-Modal."""
    return request.headers.get('x-modal') == '1'


def _formulario(request, plantilla, contexto):
    """En un modal se envía solo el formulario; si no, la página completa."""
    modal = _es_modal(request)
    contexto.update({
        'modal': modal,
        'base_plantilla': 'Profesor/modal_base.html' if modal else 'Profesor/baseProfesor.html',
    })
    return render(request, plantilla, contexto)


def _guardado(request, mensaje, destino, navegar=False):
    """Respuesta tras guardar: en un modal la página actual se recarga (o navega si se pide)."""
    messages.success(request, mensaje)
    if _es_modal(request):
        return JsonResponse({'ok': True, 'redirect': destino, 'navegar': navegar})
    return redirect(destino)


def _colores():
    """Los colores de clases y equipos con sus dos tonos, para pintar las muestras."""
    return [{'valor': valor, 'nombre': nombre, 'c1': COLORES_CLASE[valor][0], 'c2': COLORES_CLASE[valor][1]}
            for valor, nombre in Clase.COLORES]


def _clases_info():
    """Clases como tarjetas para elegir (ícono, color y cuántos alumnos tienen)."""
    clases = Clase.objects.annotate(
        total=Count('alumnos', filter=Q(alumnos__activo=True), distinct=True)
    ).order_by('-activa', 'grado', 'grupo')
    return [{'id': c.id, 'nombre': c.nombre, 'icono': c.icono, 'grado': c.grado, 'grupo': c.grupo,
             'activa': c.activa, 'total': c.total, 'c1': COLORES_CLASE.get(c.color, COLORES_CLASE['indigo'])[0],
             'c2': COLORES_CLASE.get(c.color, COLORES_CLASE['indigo'])[1]} for c in clases]


def _acomodar(objeto, hermanos, posicion, mismo_lugar=False):
    """Acomoda el objeto entre sus hermanos (al inicio, después de otro o al final) y renumera 1, 2, 3..."""
    if posicion == 'mantener' and mismo_lugar:
        return
    lista = list(hermanos.exclude(pk=objeto.pk).order_by('orden', 'id'))
    indice = len(lista)
    if posicion == 'inicio':
        indice = 0
    elif posicion.startswith('despues:'):
        referencia = posicion.split(':', 1)[1]
        indice = next((i + 1 for i, h in enumerate(lista) if str(h.id) == referencia), len(lista))
    lista.insert(indice, objeto)
    for orden, actual in enumerate(lista, start=1):
        if actual.orden != orden:
            type(actual).objects.filter(pk=actual.pk).update(orden=orden)


def _misiones_ordenadas(clase):
    return Mision.objects.filter(clase=clase).select_related('nivel', 'nivel__mundo').order_by(
        'nivel__mundo__orden', 'nivel__mundo_id', 'nivel__orden', 'nivel_id', 'orden', 'id'
    )


# ========================
# DASHBOARD
# ========================
def _barra_redondeada(x, base, ancho, alto, radio=4):
    """Path SVG de una columna con la punta redondeada y la base recta."""
    if alto <= 0:
        return ''
    r = min(radio, alto, ancho / 2)
    y = base - alto
    return (f'M{x:.1f},{base:.1f} V{y + r:.1f} Q{x:.1f},{y:.1f} {x + r:.1f},{y:.1f} '
            f'H{x + ancho - r:.1f} Q{x + ancho:.1f},{y:.1f} {x + ancho:.1f},{y + r:.1f} V{base:.1f} Z')


def _grafica_actividad(actividad):
    izq, arriba, ancho_plot, alto_plot = 34, 12, 600, 180
    base = arriba + alto_plot
    maximo = max([max(d['entregas'], d['aprobadas']) for d in actividad] + [1])
    paso = 1
    for candidato in (1, 2, 5, 10, 20, 25, 50, 100, 200, 250, 500):
        paso = candidato
        if candidato * 4 >= maximo:
            break
    tope = paso * 4
    banda = ancho_plot / len(actividad)
    barra = min(14, banda / 2 - 3)
    grupos = []
    for i, dia in enumerate(actividad):
        centro = izq + banda * i + banda / 2
        x1 = centro - barra - 1
        x2 = centro + 1
        h1 = dia['entregas'] / tope * alto_plot
        h2 = dia['aprobadas'] / tope * alto_plot
        grupos.append({
            **dia,
            'zona_x': izq + banda * i, 'zona_ancho': banda,
            'centro': centro,
            'p1': _barra_redondeada(x1, base, barra, h1),
            'p2': _barra_redondeada(x2, base, barra, h2),
            'etiqueta_visible': i % 2 == 1 or len(actividad) <= 7,
        })
    ticks = [{'valor': paso * k, 'y': base - paso * k / tope * alto_plot} for k in range(5)]
    return {'grupos': grupos, 'ticks': ticks, 'base': base, 'izq': izq, 'der': izq + ancho_plot,
            'alto': base + 22, 'ancho': izq + ancho_plot + 6, 'arriba': arriba, 'alto_plot': alto_plot,
            'total_entregas': sum(d['entregas'] for d in actividad),
            'total_aprobadas': sum(d['aprobadas'] for d in actividad)}


@login_required
def dashboard_profesor(request):
    ahora = timezone.now()
    ahora_local = timezone.localtime()
    hace_7 = ahora - timedelta(days=7)
    hoy = timezone.localdate()

    # Actividad de los últimos 14 días: entregas y misiones aprobadas
    inicio = hoy - timedelta(days=13)
    entregas = dict(
        ProgresoMision.objects.filter(fecha__date__gte=inicio).annotate(dia=TruncDate('fecha'))
        .values_list('dia').annotate(total=Count('id'))
    )
    aprobadas = dict(
        ProgresoMision.objects.filter(estado='aprobado', fecha_revision__date__gte=inicio)
        .annotate(dia=TruncDate('fecha_revision')).values_list('dia').annotate(total=Count('id'))
    )
    actividad = []
    for i in range(14):
        dia = inicio + timedelta(days=i)
        actividad.append({
            'fecha': dia.isoformat(),
            'etiqueta': dia.strftime('%d/%m'),
            'entregas': entregas.get(dia, 0),
            'aprobadas': aprobadas.get(dia, 0),
        })

    clases = list(Clase.objects.annotate(
        total_alumnos=Count('alumnos', filter=Q(alumnos__activo=True), distinct=True),
        total_misiones=Count('misiones', filter=Q(misiones__activa=True), distinct=True),
        total_mundos=Count('mundos', distinct=True),
    ).order_by('-activa', 'anio', 'grado', 'grupo'))
    for clase in clases:
        clase.misiones_pendientes = ProgresoMision.objects.filter(mision__clase=clase, estado='pendiente').count()
        posibles = clase.total_alumnos * clase.total_misiones
        hechas = ProgresoMision.objects.filter(
            mision__clase=clase, mision__activa=True, estado='aprobado', alumno__activo=True, alumno__clases=clase
        ).count()
        clase.avance = int(hechas * 100 / posibles) if posibles else 0

    top_semana = Alumno.objects.filter(activo=True).annotate(
        puntos=Coalesce(Sum('movimientos__xp', filter=Q(movimientos__fecha__gte=hace_7)), Value(0))
    ).filter(puntos__gt=0).order_by('-puntos')[:5]

    return render(request, 'Profesor/dashboard_principal.html', {
        'kpi_alumnos': Alumno.objects.filter(activo=True).count(),
        'kpi_revisiones': ProgresoMision.objects.filter(estado='pendiente').count(),
        'kpi_pedidos': CompraAlumno.objects.filter(estado='pendiente').count(),
        'kpi_xp_semana': Movimiento.objects.filter(fecha__gte=hace_7, xp__gt=0).aggregate(
            t=Coalesce(Sum('xp'), Value(0)))['t'],
        'kpi_aprobadas_semana': ProgresoMision.objects.filter(estado='aprobado', fecha_revision__gte=hace_7).count(),
        'actividad': actividad,
        'grafica': _grafica_actividad(actividad),
        'clases': clases,
        'entregas_tienda': CompraAlumno.objects.filter(estado='pendiente').select_related(
            'alumno', 'articulo').order_by('fecha_compra')[:6],
        'stock_bajo': ArticuloTienda.objects.filter(activo=True, stock__isnull=False, stock__lte=3).order_by('stock'),
        'movimientos': Movimiento.objects.select_related('alumno').order_by('-fecha')[:8],
        'top_semana': top_semana,
        'saludo': 'Buenos días' if ahora_local.hour < 12 else 'Buenas tardes' if ahora_local.hour < 19 else 'Buenas noches',
        'seccion': 'inicio',
    })


# ========================
# CLASES Y ESTRUCTURA
# ========================
@login_required
def clases_lista(request):
    clases = Clase.objects.annotate(
        total_alumnos=Count('alumnos', filter=Q(alumnos__activo=True), distinct=True),
        total_mundos=Count('mundos', distinct=True),
        total_misiones=Count('misiones', distinct=True),
    ).order_by('-activa', 'anio', 'grado', 'grupo')
    return render(request, 'Profesor/clases.html', {'clases': clases, 'seccion': 'clases'})


ICONOS_CLASE = ['🚀', '🤖', '💡', '🔬', '🧪', '🛠️', '⚙️', '💻', '🎮', '🧩', '🎨', '🌍', '📐', '📚', '🎵', '⚽', '🦖', '🌱']


@login_required
def alta_clase(request, pk=None):
    clase = get_object_or_404(Clase, pk=pk) if pk else None
    anio = timezone.localdate().year
    form = ClaseForm(request.POST or None, instance=clase, initial=None if clase else {'anio': anio})
    if request.method == 'POST' and form.is_valid():
        clase = form.save()
        return _guardado(request, f'Clase "{clase.nombre}" guardada.', reverse('clase_detalle', args=[clase.id]))
    otras = Clase.objects.exclude(pk=clase.pk) if clase else Clase.objects.all()
    return _formulario(request, 'Profesor/clase_form.html', {
        'form': form,
        'clase': clase,
        'iconos': ICONOS_CLASE,
        'colores': _colores(),
        'grados': ['1°', '2°', '3°', '4°', '5°', '6°'],
        'grupos': ['A', 'B', 'C', 'D'],
        'anios': [anio - 1, anio, anio + 1],
        'existentes': [{'nombre': c.nombre, 'grado': c.grado, 'grupo': c.grupo, 'anio': c.anio} for c in otras],
        'total_alumnos': clase.alumnos.filter(activo=True).count() if clase else 0,
        'total_mundos': clase.mundos.count() if clase else 0,
        'seccion': 'clases',
    })


@login_required
def clase_detalle(request, clase_id):
    clase = get_object_or_404(Clase, id=clase_id)
    alumnos_ids = list(clase.alumnos.filter(activo=True).values_list('id', flat=True))
    stats = {
        fila['mision_id']: fila for fila in ProgresoMision.objects.filter(
            mision__clase=clase, alumno_id__in=alumnos_ids
        ).values('mision_id').annotate(
            aprobadas=Count('id', filter=Q(estado='aprobado')),
            pendientes=Count('id', filter=Q(estado='pendiente')),
        )
    }

    def armar_nivel(nivel):
        misiones = []
        for mision in nivel.misiones.all():
            datos = stats.get(mision.id, {})
            mision.aprobadas = datos.get('aprobadas', 0)
            mision.pendientes = datos.get('pendientes', 0)
            mision.info = juego.TIPO_INFO.get(mision.tipo)
            misiones.append(mision)
        misiones.sort(key=lambda m: (m.orden, m.id))
        return {'nivel': nivel, 'misiones': misiones}

    mundos = []
    for mundo in clase.mundos.prefetch_related('niveles__misiones'):
        mundos.append({
            'mundo': mundo,
            'tema': juego.TEMA_INFO.get(mundo.tema),
            'niveles': [armar_nivel(n) for n in mundo.niveles.all()],
        })
    sueltos = [armar_nivel(n) for n in clase.niveles.filter(mundo__isnull=True).prefetch_related('misiones')]

    return render(request, 'Profesor/clase_detalle.html', {
        'clase': clase,
        'mundos': mundos,
        'sueltos': sueltos,
        'total_alumnos': len(alumnos_ids),
        'pendientes': ProgresoMision.objects.filter(mision__clase=clase, estado='pendiente').count(),
        'seccion': 'clases',
    })


NOMBRES_MUNDO = {
    'bosque': ['Bosque de los Circuitos', 'Selva de los Sensores', 'Bosque Encantado'],
    'oceano': ['Océano de los Datos', 'Arrecife Robótico', 'Abismo de los Bits'],
    'desierto': ['Desierto de los Engranes', 'Dunas del Código', 'Oasis Digital'],
    'volcan': ['Volcán de los Motores', 'Montaña de Fuego', 'Cráter de la Energía'],
    'hielo': ['Reino de Hielo', 'Glaciar de los Algoritmos', 'Polo Norte Digital'],
    'espacio': ['Galaxia Arduino', 'Estación Espacial', 'Planeta de los Robots'],
    'ciudad': ['Ciudad Neón', 'Metrópolis Robótica', 'Ciudad de los Inventos'],
    'dulces': ['Tierra de Dulces', 'Fábrica de Golosinas', 'Reino de Caramelo'],
}


@login_required
def mundo_form(request, pk=None):
    mundo = get_object_or_404(Mundo, pk=pk) if pk else None
    inicial = {}
    if not mundo and request.GET.get('clase'):
        clase = Clase.objects.filter(pk=request.GET['clase']).first()
        if clase:
            siguiente = clase.mundos.count() + 1
            tema = [t[0] for t in Mundo.TEMAS][(siguiente - 1) % len(Mundo.TEMAS)]
            usados = set(clase.mundos.values_list('nombre', flat=True))
            nombre = next((n for n in NOMBRES_MUNDO[tema] if n not in usados), f'Mundo {siguiente}')
            inicial = {'clase': clase, 'tema': tema, 'nombre': nombre}
    form = MundoForm(request.POST or None, instance=mundo, initial=inicial)
    clase_anterior = mundo.clase_id if mundo else None
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            mundo = form.save()
            _acomodar(mundo, Mundo.objects.filter(clase=mundo.clase), request.POST.get('posicion', 'final'),
                      clase_anterior == mundo.clase_id)
        return _guardado(request, f'Mundo "{mundo.nombre}" guardado.', reverse('clase_detalle', args=[mundo.clase_id]))
    temas = {clave: {**info, 'nombres': NOMBRES_MUNDO.get(clave, [])} for clave, info in juego.TEMA_INFO.items()}
    return _formulario(request, 'Profesor/mundo_form.html', {
        'form': form,
        'mundo': mundo,
        'temas': juego.TEMA_INFO,
        'temas_json': temas,
        'mundos_por_clase': {
            c.id: [{'id': m.id, 'nombre': m.nombre, 'emoji': juego.TEMA_INFO.get(m.tema, {}).get('emoji', '🌍')}
                   for m in c.mundos.all()]
            for c in Clase.objects.prefetch_related('mundos')
        },
        'seccion': 'clases',
    })


@login_required
def alta_nivel(request, pk=None):
    nivel = get_object_or_404(Nivel, pk=pk) if pk else None
    inicial = {}
    if not nivel and request.GET.get('mundo'):
        mundo = Mundo.objects.filter(pk=request.GET['mundo']).first()
        if mundo:
            inicial = {'mundo': mundo.id, 'nombre': f'Nivel {mundo.niveles.count() + 1}'}
    form = NivelForm(request.POST or None, instance=nivel, initial=inicial)
    if nivel and not request.POST:
        form.initial['mundo'] = nivel.mundo_id
    mundo_anterior = nivel.mundo_id if nivel else None
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            nivel = form.save()
            _acomodar(nivel, Nivel.objects.filter(mundo=nivel.mundo), request.POST.get('posicion', 'final'),
                      mundo_anterior == nivel.mundo_id)
        return _guardado(request, f'Nivel "{nivel.nombre}" guardado.', reverse('clase_detalle', args=[nivel.clase_id]))
    mundos = {}
    for mundo in Mundo.objects.select_related('clase').prefetch_related('niveles').annotate(n=Count('niveles')):
        mundos[mundo.id] = {
            'nombre': mundo.nombre, 'tema': mundo.tema,
            'emoji': juego.TEMA_INFO.get(mundo.tema, {}).get('emoji', '🌍'),
            'clase': f'{mundo.clase.icono} {mundo.clase.nombre} · {mundo.clase.grado} {mundo.clase.grupo}',
            'niveles': [{'id': n.id, 'nombre': n.nombre} for n in mundo.niveles.all()],
        }
    return _formulario(request, 'Profesor/nivel_form.html', {
        'form': form,
        'nivel': nivel,
        'mundos': mundos,
        'misiones_nivel': nivel.misiones.count() if nivel else 0,
        'seccion': 'clases',
    })


def _aplicar_posicion(mision, posicion, nivel_anterior=None):
    """Acomoda la misión dentro de su nivel y renumera el orden 1, 2, 3..."""
    _acomodar(mision, Mision.objects.filter(nivel=mision.nivel), posicion, nivel_anterior == mision.nivel_id)


def _estructura_niveles():
    """Clases → mundos → niveles → misiones para elegir dónde va una misión."""
    clases = []
    for clase in Clase.objects.order_by('-activa', 'grado', 'grupo').prefetch_related(
            'mundos__niveles__misiones', 'niveles__misiones'):
        def nivel_json(nivel):
            misiones = sorted(nivel.misiones.all(), key=lambda m: (m.orden, m.id))
            return {'id': nivel.id, 'nombre': nivel.nombre,
                    'misiones': [{'id': m.id, 'nombre': m.nombre} for m in misiones]}
        mundos = [{
            'id': mundo.id, 'nombre': mundo.nombre, 'emoji': juego.TEMA_INFO.get(mundo.tema, {}).get('emoji', '🌍'),
            'niveles': [nivel_json(n) for n in mundo.niveles.all()],
        } for mundo in clase.mundos.all()]
        sueltos = [n for n in clase.niveles.all() if n.mundo_id is None]
        if sueltos:
            mundos.append({'id': 0, 'nombre': 'Niveles sin mundo', 'emoji': '⚠️', 'niveles': [nivel_json(n) for n in sueltos]})
        clases.append({'id': clase.id, 'nombre': f'{clase.icono} {clase.nombre} · {clase.grado} {clase.grupo}',
                       'mundos': mundos})
    return clases


def _guardar_preguntas(mision, datos):
    """Crea, actualiza o borra preguntas del cuestionario según el JSON del constructor."""
    existentes = {p.id: p for p in mision.preguntas.prefetch_related('opciones')}
    conservar = set()
    for orden, dato in enumerate(datos, start=1):
        texto = (dato.get('texto') or '').strip()
        if not texto:
            continue
        pregunta = existentes.get(dato.get('id'))
        puntos = max(1, int(dato.get('puntos') or 10))
        if pregunta:
            pregunta.texto, pregunta.orden, pregunta.puntos = texto, orden, puntos
            pregunta.save()
        else:
            pregunta = Pregunta.objects.create(mision=mision, texto=texto, orden=orden, puntos=puntos)
        conservar.add(pregunta.id)

        opciones = {o.id: o for o in pregunta.opciones.all()}
        conservar_opc = set()
        for dato_opc in dato.get('opciones', []):
            texto_opc = (dato_opc.get('texto') or '').strip()
            if not texto_opc:
                continue
            opcion = opciones.get(dato_opc.get('id'))
            correcta = bool(dato_opc.get('es_correcta'))
            if opcion:
                opcion.texto, opcion.es_correcta = texto_opc, correcta
                opcion.save()
            else:
                opcion = OpcionRespuesta.objects.create(pregunta=pregunta, texto=texto_opc, es_correcta=correcta)
            conservar_opc.add(opcion.id)
        pregunta.opciones.exclude(id__in=conservar_opc).delete()
    mision.preguntas.exclude(id__in=conservar).delete()


@login_required
def alta_mision(request, pk=None):
    mision = get_object_or_404(Mision, pk=pk) if pk else None
    inicial = {}
    if not mision and request.GET.get('nivel'):
        nivel = Nivel.objects.filter(pk=request.GET['nivel']).first()
        if nivel:
            inicial = {'nivel': nivel.id}
    if not mision:
        inicial.update({'xp_recompensa': 50, 'monedas_recompensa': 15, 'activa': True})
        # «Agregar al mapa» desde la biblioteca de juegos
        juego_elegido = Juego.objects.filter(pk=request.GET.get('juego'), activo=True).first()             if (request.GET.get('juego') or '').isdigit() else None
        if juego_elegido:
            inicial.update({'tipo': 'juego', 'juego': juego_elegido.id, 'nombre': juego_elegido.nombre,
                            'descripcion': juego_elegido.descripcion or minijuegos.PLANTILLAS_INFO[juego_elegido.plantilla]['como']})
    form = MisionForm(request.POST or None, request.FILES or None, instance=mision, initial=inicial)
    if mision and not request.POST:
        form.initial['nivel'] = mision.nivel_id

    error_cuestionario = ''
    nivel_anterior = mision.nivel_id if mision else None
    if request.method == 'POST' and form.is_valid():
        try:
            with transaction.atomic():
                guardada = form.save()
                _aplicar_posicion(guardada, request.POST.get('posicion', 'final'), nivel_anterior)
                if guardada.tipo == 'cuestionario':
                    _guardar_preguntas(guardada, json.loads(request.POST.get('preguntas_json') or '[]'))
            return _guardado(request, f'Misión "{guardada.nombre}" guardada.',
                             _siguiente(request, reverse('clase_detalle', args=[guardada.clase_id])))
        except (ValueError, TypeError) as error:
            error_cuestionario = f'Ocurrió un error al procesar el cuestionario: {error}'

    preguntas = []
    if request.method == 'POST':
        # Si el formulario tuvo errores, se conservan las preguntas que el profe ya escribió
        try:
            preguntas = json.loads(request.POST.get('preguntas_json') or '[]')
        except ValueError:
            preguntas = []
    elif mision:
        for pregunta in mision.preguntas.prefetch_related('opciones'):
            preguntas.append({
                'id': pregunta.id, 'texto': pregunta.texto, 'puntos': pregunta.puntos,
                'opciones': [{'id': o.id, 'texto': o.texto, 'es_correcta': o.es_correcta}
                             for o in pregunta.opciones.all()],
            })
    return _formulario(request, 'Administrativo/alta_mision.html', {
        'form': form,
        'mision': mision,
        'preguntas': preguntas,
        'tipos': juego.TIPO_INFO,
        'juegos_data': {j.id: {'nombre': j.nombre, 'icono': j.icono, 'meta': minijuegos.descripcion_meta(j),
                               'descripcion': j.descripcion or minijuegos.PLANTILLAS_INFO[j.plantilla]['como']}
                        for j in Juego.objects.all()},
        'estructura': _estructura_niveles(),
        'posicion': request.POST.get('posicion', 'mantener' if mision else 'final'),
        'error_cuestionario': error_cuestionario,
        'seccion': 'clases',
    })


@require_POST
@login_required
def duplicar_mision(request, pk):
    original = get_object_or_404(Mision, pk=pk)
    with transaction.atomic():
        preguntas = list(original.preguntas.prefetch_related('opciones'))
        copia = original
        copia.pk = None
        copia.nombre = f'{original.nombre} (copia)'[:200]
        copia.activa = False
        copia.orden = original.orden + 1
        copia.save()
        for pregunta in preguntas:
            opciones = list(pregunta.opciones.all())
            pregunta.pk = None
            pregunta.mision = copia
            pregunta.save()
            for opcion in opciones:
                opcion.pk = None
                opcion.pregunta = pregunta
                opcion.save()
    messages.success(request, 'Misión duplicada. La copia está oculta para que la edites antes de publicarla.')
    editar = reverse('editar_mision', args=[copia.pk])
    return redirect(f"{reverse('clase_detalle', args=[copia.clase_id])}?abrir={editar}")


CONDICIONES_INFO = [
    # (valor, emoji, título corto, ejemplo)
    ('misiones', '🏁', 'Completar misiones', 'Ej. completa 10 misiones'),
    ('xp', '⚡', 'Juntar XP', 'Ej. llega a 1,000 XP'),
    ('nivel', '📈', 'Subir de nivel', 'Ej. llega al nivel 5'),
    ('racha', '🔥', 'Racha de días', 'Ej. entra 7 días seguidos'),
    ('perfectos', '💯', 'Sacar 10', 'Ej. 10 en 3 cuestionarios'),
    ('jefes', '👾', 'Vencer jefes', 'Ej. derrota 2 jefes'),
    ('mundos', '🌍', 'Completar mundos', 'Ej. termina un mundo'),
    ('diarias', '📅', 'Misiones del día', 'Ej. reclama 10 diarias'),
    ('monedas', '💰', 'Ganar monedas', 'Ej. gana 500 monedas'),
    ('compras', '🛒', 'Comprar en la tienda', 'Ej. haz 3 compras'),
    ('juegos', '🎮', 'Jugar minijuegos', 'Ej. juega 20 partidas'),
    ('ppm', '⌨️', 'Escribir rápido', 'Ej. 30 palabras por minuto'),
    ('manual', '🎁', 'Lo das tú', 'Reconocimiento especial'),
]

TIPOS_ARTICULO = [
    # (valor, emoji, título corto, descripción)
    ('fisico', '🎁', 'Premio físico', 'Algo que entregas en mano'),
    ('permiso', '🎟️', 'Privilegio en clase', 'Un permiso o beneficio'),
    ('avatar', '🤖', 'Avatar', 'Imagen para su perfil'),
    ('otro', '💎', 'Otro digital', 'Algo dentro de la plataforma'),
]

OBJETIVOS_DIARIA = [
    # (valor, emoji, título corto, cómo se cuenta)
    ('login', '👋', 'Entrar', 'Se cumple al abrir su mapa'),
    ('misiones', '🗺️', 'Completar misiones', 'Misiones del mapa enviadas'),
    ('cuestionarios', '🧠', 'Responder cuestionarios', 'Cuestionarios terminados'),
    ('perfecto', '💯', 'Sacar 10', 'Cuestionarios sin errores'),
    ('xp', '🔋', 'Ganar XP', 'Toda la XP del día'),
    ('evidencias', '🛠️', 'Enviar trabajos', 'Entregas para revisión'),
    ('tienda', '🛍️', 'Visitar la tienda', 'Solo entrar, sin comprar'),
    ('juegos', '🎮', 'Jugar minijuegos', 'Partidas del mapa o del Arcade'),
]

ICONOS_DIARIA = ['👋', '🗺️', '⚡', '🧠', '💯', '🔋', '🛠️', '🛍️', '⭐', '🔥', '🎯', '🚀', '🤖', '💡', '📚', '🏃',
                 '🌞', '🎮', '🧩', '✏️', '🎨', '🔬', '🏆', '🍀']

ICONOS_LOGRO = ['🏅', '🏆', '⭐', '🌟', '✨', '🔥', '💯', '🧠', '🎓', '🚀', '🤖', '💡', '⚙️', '🛠️', '🎯', '👑',
                '💎', '🐉', '🦄', '🌈', '🎨', '📚', '🧭', '⚡', '💰', '🎁', '👣', '🦸', '🏰', '🛡️', '🐢', '🦉']

MODELOS_ELIMINABLES = {
    'clase': (Clase, 'clases_lista'),
    'mundo': (Mundo, None),
    'nivel': (Nivel, None),
    'mision': (Mision, None),
    'alumno': (Alumno, 'alumnos_lista'),
    'articulo': (ArticuloTienda, 'tienda_admin'),
    'diaria': (MisionDiaria, 'diarias_lista'),
    'logro': (Logro, 'logros_lista'),
    'equipo': (Equipo, 'equipos_lista'),
    'meta': (MetaGrupo, 'metas_lista'),
    'comportamiento': (Comportamiento, 'comportamientos_lista'),
    'juego': (Juego, 'juegos_lista'),
}


@require_POST
@login_required
def eliminar(request, modelo, pk):
    if modelo not in MODELOS_ELIMINABLES:
        return HttpResponse(status=404)
    clase_modelo, destino = MODELOS_ELIMINABLES[modelo]
    objeto = get_object_or_404(clase_modelo, pk=pk)
    nombre = str(getattr(objeto, 'nombre', None) or getattr(objeto, 'titulo', None) or objeto)
    por_defecto = reverse(destino) if destino else reverse('clase_detalle', args=[objeto.clase_id])
    objeto.delete()
    messages.success(request, f'"{nombre}" se eliminó.')
    return redirect(_siguiente(request, por_defecto))


# ========================
# REVISIONES Y CALIFICACIONES
# ========================
@login_required
def revisiones(request):
    clase_id = request.GET.get('clase')
    tipo = request.GET.get('tipo')
    entregas = ProgresoMision.objects.filter(estado='pendiente').select_related(
        'alumno', 'mision', 'mision__clase', 'mision__nivel', 'mision__nivel__mundo'
    ).order_by('fecha')
    clase = None
    if clase_id and clase_id.isdigit():
        clase = Clase.objects.filter(pk=clase_id).first()
        entregas = entregas.filter(mision__clase_id=clase_id)
    if tipo:
        entregas = entregas.filter(mision__tipo=tipo)

    clases = Clase.objects.annotate(
        pendientes=Count('misiones__progresomision', filter=Q(misiones__progresomision__estado='pendiente'))
    ).order_by('grado', 'grupo')
    recientes = ProgresoMision.objects.filter(revisado=True, fecha_revision__isnull=False).select_related(
        'alumno', 'mision').order_by('-fecha_revision')[:8]
    return render(request, 'Profesor/panel_calificar.html', {
        'entregas': entregas,
        'clases': clases,
        'clase': clase,
        'tipo': tipo,
        'tipos': Mision.TIPO_MISION,
        'recientes': recientes,
        'seccion': 'revisiones',
    })


@login_required
def panel_profesor(request, clase_id):
    return redirect(f"{reverse('revisiones')}?clase={clase_id}")


def _leer_calificacion(valor):
    """Acepta calificación de 0 a 10 (con decimales) y la convierte a 0-100."""
    try:
        numero = float(str(valor).replace(',', '.'))
    except (TypeError, ValueError):
        return None
    if numero > 10:
        numero = numero / 10
    return int(round(max(0, min(10, numero)) * 10))


@require_POST
@login_required
def aprobar_mision(request, progreso_id):
    progreso = get_object_or_404(ProgresoMision.objects.select_related('alumno', 'mision'), id=progreso_id)
    puntuacion = _leer_calificacion(request.POST.get('calificacion', 10))
    if puntuacion is None:
        messages.error(request, 'La calificación no es válida.')
    else:
        monedas = juego.calificar(progreso, puntuacion, request.POST.get('comentario', '').strip())
        messages.success(request, f'Misión de {progreso.alumno.nombre} aprobada con {puntuacion / 10:.1f} '
                                  f'(+{monedas} monedas).')
    return redirect(_siguiente(request, reverse('revisiones')))


@require_POST
@login_required
def rechazar_mision(request, progreso_id):
    progreso = get_object_or_404(ProgresoMision.objects.select_related('alumno', 'mision'), id=progreso_id)
    if progreso.estado == 'pendiente':
        juego.rechazar(progreso, request.POST.get('comentario', '').strip())
        messages.warning(request, f'Misión de {progreso.alumno.nombre} devuelta para corrección.')
    return redirect(_siguiente(request, reverse('revisiones')))


@login_required
def calificar_mision(request, mision_id):
    mision = get_object_or_404(Mision.objects.select_related('clase', 'nivel', 'nivel__mundo'), id=mision_id)
    alumnos = list(mision.clase.alumnos.filter(activo=True).order_by('apellido', 'nombre'))
    progresos = {p.alumno_id: p for p in ProgresoMision.objects.filter(mision=mision)}

    if request.method == 'POST':
        comentario = request.POST.get('comentario', '').strip() or None
        calificados = 0
        for alumno in alumnos:
            valor = request.POST.get(f'calif_{alumno.id}', '').strip()
            if not valor:
                continue
            puntuacion = _leer_calificacion(valor)
            if puntuacion is None:
                continue
            progreso = progresos.get(alumno.id) or ProgresoMision(alumno=alumno, mision=mision)
            if progreso.estado == 'aprobado' and progreso.puntuacion == puntuacion:
                continue
            juego.calificar(progreso, puntuacion, comentario)
            calificados += 1
        return _guardado(request, f'{calificados} calificaciones guardadas para "{mision.nombre}".',
                         _siguiente(request, reverse('calificar_mision', args=[mision.id])))

    for alumno in alumnos:
        alumno.progreso = progresos.get(alumno.id)
    return _formulario(request, 'Profesor/calificar_mision.html', {
        'mision': mision,
        'alumnos': alumnos,
        'info': juego.TIPO_INFO.get(mision.tipo),
        'seccion': 'revisiones',
    })


@login_required
def calificaciones_inicio(request):
    clase = Clase.objects.filter(activa=True).order_by('grado', 'grupo').first() or Clase.objects.first()
    if not clase:
        return redirect('clases_lista')
    return redirect('matriz_calificaciones', clase_id=clase.id)


def _datos_matriz(clase):
    alumnos = list(clase.alumnos.filter(activo=True).order_by('apellido', 'nombre'))
    misiones = list(_misiones_ordenadas(clase))
    progresos = {
        (p.alumno_id, p.mision_id): p
        for p in ProgresoMision.objects.filter(alumno__in=alumnos, mision__in=misiones)
    }
    matriz = []
    for alumno in alumnos:
        celdas = [{'mision': m, 'progreso': progresos.get((alumno.id, m.id))} for m in misiones]
        notas = [c['progreso'].puntuacion for c in celdas if c['progreso'] and c['progreso'].estado == 'aprobado']
        matriz.append({
            'alumno': alumno,
            'calificaciones': celdas,
            'promedio': sum(notas) / len(notas) / 10 if notas else None,
            'completadas': len(notas),
        })
    for indice, mision in enumerate(misiones):
        aprobadas = [f['calificaciones'][indice]['progreso'] for f in matriz
                     if f['calificaciones'][indice]['progreso']
                     and f['calificaciones'][indice]['progreso'].estado == 'aprobado']
        mision.completado = int(len(aprobadas) * 100 / len(alumnos)) if alumnos else 0
        mision.promedio = sum(p.puntuacion for p in aprobadas) / len(aprobadas) / 10 if aprobadas else None
    return alumnos, misiones, matriz


@login_required
def matriz_calificaciones(request, clase_id):
    clase = get_object_or_404(Clase, id=clase_id)
    alumnos, misiones, matriz = _datos_matriz(clase)
    promedios = [f['promedio'] for f in matriz if f['promedio'] is not None]
    return render(request, 'Profesor/matriz_calificaciones.html', {
        'clase': clase,
        'misiones_columnas': misiones,
        'matriz': matriz,
        'promedio_grupo': sum(promedios) / len(promedios) if promedios else None,
        'clases': Clase.objects.order_by('grado', 'grupo'),
        'seccion': 'calificaciones',
    })


@login_required
def exportar_matriz_excel(request, clase_id):
    clase = get_object_or_404(Clase, id=clase_id)
    alumnos, misiones, matriz = _datos_matriz(clase)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Matriz {clase.grado} {clase.grupo}"[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="142C45", end_color="142C45", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center")

    headers = ['Alumno'] + [
        f"{m.nombre}\n({m.nivel.mundo.nombre + ' · ' if m.nivel.mundo else ''}{m.nivel.nombre})" for m in misiones
    ] + ['Promedio', 'XP', 'Monedas']
    ws.append(headers)
    for col_num, cell in enumerate(ws[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        ws.column_dimensions[get_column_letter(col_num)].width = 35 if col_num == 1 else 18
    ws.row_dimensions[1].height = 45

    for fila in matriz:
        alumno = fila['alumno']
        valores = [f"{alumno.apellido}, {alumno.nombre}"]
        for celda in fila['calificaciones']:
            progreso = celda['progreso']
            if not progreso:
                valores.append("Sin entrega")
            elif progreso.estado == 'aprobado':
                valores.append(progreso.calificacion_escala_10)
            else:
                valores.append(progreso.get_estado_display())
        valores += [round(fila['promedio'], 1) if fila['promedio'] is not None else '-',
                    alumno.xp_total, alumno.monedas]
        ws.append(valores)

    for row in ws.iter_rows(min_row=2, max_col=len(headers), max_row=len(alumnos) + 1):
        row[0].alignment = left
        for cell in row[1:]:
            cell.alignment = center
    ws.freeze_panes = 'B2'

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    nombre_archivo = f"Calificaciones_{clase.nombre}_{clase.grado}_{clase.grupo}.xlsx".replace(" ", "_")
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    wb.save(response)
    return response


# ========================
# ALUMNOS
# ========================
PINES_FACILES = {'1234', '4321', '0123', '2580', '1212', '1122', '9876', '6789', '0987', '1004'}


def _pin_al_azar():
    """PIN de 4 dígitos que no sea fácil de adivinar (ni 1111 ni 1234)."""
    while True:
        pin = f'{random.randint(0, 9999):04d}'
        if len(set(pin)) > 1 and pin not in PINES_FACILES:
            return pin


@login_required
def alumnos_lista(request):
    clase_id = request.GET.get('clase')
    alumnos = Alumno.objects.prefetch_related('clases').annotate(
        aprobadas=Count('progresomision', filter=Q(progresomision__estado='aprobado'), distinct=True)
    ).order_by('-activo', 'apellido', 'nombre')
    clase = None
    if clase_id and clase_id.isdigit():
        clase = Clase.objects.filter(pk=clase_id).first()
        alumnos = alumnos.filter(clases__id=clase_id)
    return render(request, 'Profesor/alumnos.html', {
        'alumnos': alumnos.distinct(),
        'clases': Clase.objects.order_by('grado', 'grupo'),
        'clase': clase,
        'seccion': 'alumnos',
    })


@login_required
def alta_alumno(request, pk=None):
    alumno = get_object_or_404(Alumno, pk=pk) if pk else None
    inicial = {} if alumno else {'pin': _pin_al_azar(), 'activo': True}
    if not alumno and request.GET.get('clase'):
        inicial['clases'] = [request.GET['clase']]
    form = AlumnoForm(request.POST or None, request.FILES or None, instance=alumno, initial=inicial)
    if request.method == 'POST' and form.is_valid():
        alumno = form.save()
        return _guardado(request, f'{alumno.nombre} {alumno.apellido} quedó registrado con PIN {alumno.pin}.',
                         reverse('alumno_detalle', args=[alumno.pk]))
    return _formulario(request, 'Profesor/alumno_form.html', {
        'form': form,
        'alumno': alumno,
        'clases': _clases_info(),
        'seleccionadas': {str(v) for v in (form['clases'].value() or [])},
        'avatar': avatar_url(alumno) if alumno else '',
        # El robot automático que tendrá (o ya tiene) si no sube foto
        'robot': avatar_url(Alumno(pk=alumno.pk)) if alumno else avatar_url(Alumno(pk=0)),
        'tiene_foto': bool(alumno and alumno.avatar),
        'seccion': 'alumnos',
    })


@login_required
def alta_masiva(request):
    form = AltaMasivaForm(request.POST or None, initial={'clase': request.GET.get('clase')})
    if request.method == 'POST' and form.is_valid():
        clase = form.cleaned_data['clase']
        creados = []
        with transaction.atomic():
            for linea in form.cleaned_data['nombres'].splitlines():
                nombre, apellido = separar_nombre(linea)
                if not nombre:
                    continue
                alumno = Alumno.objects.create(
                    nombre=nombre, apellido=apellido,
                    pin=_pin_al_azar() if form.cleaned_data['pin_aleatorio'] else '0000',
                )
                alumno.clases.add(clase)
                creados.append(str(alumno.id))
        return _guardado(request, f'Se registraron {len(creados)} alumnos en {clase.nombre}. ¡Imprime sus credenciales!',
                         f"{reverse('credenciales')}?ids={','.join(creados)}", navegar=True)
    nombres_por_clase = {}
    for alumno in Alumno.objects.filter(activo=True).prefetch_related('clases'):
        for clase in alumno.clases.all():
            nombres_por_clase.setdefault(clase.id, []).append(f'{alumno.nombre} {alumno.apellido}')
    return _formulario(request, 'Profesor/alta_masiva.html', {
        'form': form,
        'clases': _clases_info(),
        'clase_elegida': str(form['clase'].value() or ''),
        'nombres_por_clase': nombres_por_clase,
        'seccion': 'alumnos',
    })


@login_required
def credenciales(request):
    alumnos = Alumno.objects.filter(activo=True).prefetch_related('clases').order_by('apellido', 'nombre')
    clase = None
    if request.GET.get('ids'):
        ids = [i for i in request.GET['ids'].split(',') if i.isdigit()]
        alumnos = alumnos.filter(id__in=ids)
    elif request.GET.get('clase', '').isdigit():
        clase = Clase.objects.filter(pk=request.GET['clase']).first()
        alumnos = alumnos.filter(clases__id=request.GET['clase'])
    return render(request, 'Profesor/credenciales.html', {
        'alumnos': alumnos.distinct(),
        'clase': clase,
        'url_sitio': request.build_absolute_uri('/'),
    })


@login_required
def alumno_detalle(request, pk):
    alumno = get_object_or_404(Alumno.objects.prefetch_related('clases'), pk=pk)
    avance_clases = []
    for clase in alumno.clases.all():
        ruta = juego.construir_ruta(alumno, clase)
        avance_clases.append({'clase': clase, 'ruta': ruta})
    return render(request, 'Profesor/alumno_detalle.html', {
        'alumno': alumno,
        'stats': juego.estadisticas_alumno(alumno),
        'avance_clases': avance_clases,
        'progresos': ProgresoMision.objects.filter(alumno=alumno).select_related(
            'mision', 'mision__clase').order_by('-fecha')[:30],
        'compras': alumno.compras.select_related('articulo').order_by('-fecha_compra')[:20],
        'movimientos': alumno.movimientos.all()[:30],
        'logros': alumno.logros.select_related('logro'),
        'logros_disponibles': Logro.objects.filter(activo=True).exclude(alumnos__alumno=alumno),
        'seccion': 'alumnos',
    })


@require_POST
@login_required
def alumno_ajuste(request, pk):
    alumno = get_object_or_404(Alumno, pk=pk)
    motivo = request.POST.get('motivo', '').strip()
    try:
        monedas = int(request.POST.get('monedas') or 0)
        xp = int(request.POST.get('xp') or 0)
    except ValueError:
        messages.error(request, 'Escribe cantidades válidas.')
        return redirect('alumno_detalle', pk=pk)

    if monedas:
        juego.ajustar_monedas(alumno, monedas, motivo or 'Ajuste del profesor')
    if xp:
        xp = max(xp, -alumno.xp_total)
        juego.otorgar(alumno, xp=xp, tipo='ajuste', descripcion=motivo or 'Ajuste de XP del profesor')
        if xp > 0:
            juego.notificar(alumno, 'monedas', f'¡El profe te dio +{xp} XP!', motivo, '⚡')
    if monedas or xp:
        messages.success(request, f'Se actualizó a {alumno.nombre}.')
    return redirect('alumno_detalle', pk=pk)


@require_POST
@login_required
def otorgar_logro(request):
    logro = get_object_or_404(Logro, pk=request.POST.get('logro'))
    ids = request.POST.getlist('alumnos')
    otorgados = sum(1 for alumno in Alumno.objects.filter(id__in=ids) if juego.otorgar_logro(alumno, logro))
    messages.success(request, f'"{logro.nombre}" se otorgó a {otorgados} alumno(s). 🏅')
    return redirect(_siguiente(request, reverse('logros_lista')))


# ========================
# TIENDA Y PEDIDOS
# ========================
@login_required
def tienda_admin(request):
    articulos = ArticuloTienda.objects.annotate(
        vendidos=Count('compras', filter=~Q(compras__estado='cancelado')),
        por_entregar=Count('compras', filter=Q(compras__estado='pendiente')),
    ).order_by('-activo', '-destacado', 'nombre')
    return render(request, 'Profesor/tienda_admin.html', {
        'articulos': articulos,
        'seccion': 'tienda',
    })


@login_required
def alta_articulo(request, pk=None):
    articulo = get_object_or_404(ArticuloTienda, pk=pk) if pk else None
    inicial = {} if articulo else {'tipo': 'fisico', 'costo_monedas': 50, 'stock': 10, 'nivel_requerido': 1,
                                   'requiere_validacion': True, 'activo': True}
    form = ArticuloTiendaForm(request.POST or None, request.FILES or None, instance=articulo, initial=inicial)
    if request.method == 'POST' and form.is_valid():
        articulo = form.save()
        return _guardado(request, f'"{articulo.nombre}" guardado en la tienda.', reverse('tienda_admin'))
    return _formulario(request, 'Profesor/articulo_form.html', {
        'form': form,
        'articulo': articulo,
        'tipos': TIPOS_ARTICULO,
        'monedas_alumnos': list(Alumno.objects.filter(activo=True).values_list('monedas', flat=True)),
        'rangos': juego.info_rangos(),
        'seccion': 'tienda',
    })


@require_POST
@login_required
def ajustar_stock(request, pk):
    articulo = get_object_or_404(ArticuloTienda, pk=pk)
    accion = request.POST.get('accion')
    try:
        cantidad = int(request.POST.get('cantidad') or 1)
    except ValueError:
        cantidad = 1
    if accion == 'ilimitado':
        articulo.stock = None
    elif accion == 'fijar':
        articulo.stock = max(0, cantidad)
    else:
        actual = articulo.stock or 0
        articulo.stock = max(0, actual + (cantidad if accion == 'sumar' else -cantidad))
    articulo.save(update_fields=['stock'])
    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({'ok': True, 'stock': articulo.stock})
    messages.success(request, f'Stock de "{articulo.nombre}" actualizado.')
    return redirect(_siguiente(request, reverse('tienda_admin')))


@require_POST
@login_required
def alternar_activo(request, modelo, pk):
    modelos = {'articulo': ArticuloTienda, 'diaria': MisionDiaria, 'logro': Logro, 'mision': Mision,
               'clase': Clase, 'alumno': Alumno, 'meta': MetaGrupo, 'comportamiento': Comportamiento, 'juego': Juego}
    if modelo not in modelos:
        return HttpResponse(status=404)
    objeto = get_object_or_404(modelos[modelo], pk=pk)
    campo = 'activa' if hasattr(objeto, 'activa') else 'activo'
    setattr(objeto, campo, not getattr(objeto, campo))
    objeto.save(update_fields=[campo])
    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({'ok': True, 'activo': getattr(objeto, campo)})
    return redirect(_siguiente(request, reverse('dashboard_profesor')))


@login_required
def pedidos(request):
    estado = request.GET.get('estado', 'pendiente')
    if estado not in dict(CompraAlumno.ESTADOS):
        estado = 'pendiente'
    compras = CompraAlumno.objects.filter(estado=estado).select_related('alumno', 'articulo').prefetch_related(
        'alumno__clases')
    compras = compras.order_by('fecha_compra' if estado == 'pendiente' else '-fecha_compra')
    conteos = dict(CompraAlumno.objects.values_list('estado').annotate(t=Count('id')))
    pagina = Paginator(compras, 40).get_page(request.GET.get('page'))
    return render(request, 'Profesor/pedidos.html', {
        'pagina': pagina,
        'estado': estado,
        'conteos': conteos,
        'seccion': 'pedidos',
    })


@require_POST
@login_required
def entregar_articulo_tienda(request, compra_id):
    compra = get_object_or_404(CompraAlumno.objects.select_related('alumno', 'articulo'), id=compra_id)
    if juego.entregar_compra(compra):
        messages.success(request, f'¡"{compra.articulo.nombre}" entregado a {compra.alumno.nombre}!')
    return redirect(_siguiente(request, reverse('pedidos')))


@require_POST
@login_required
def cancelar_compra(request, compra_id):
    compra = get_object_or_404(CompraAlumno.objects.select_related('alumno', 'articulo'), id=compra_id)
    if juego.cancelar_compra(compra):
        messages.success(request, f'Compra cancelada: se devolvieron {compra.costo_pagado} monedas a {compra.alumno.nombre}.')
    return redirect(_siguiente(request, reverse('pedidos')))


# ========================
# ECONOMÍA
# ========================
@login_required
def gestionar_monedas(request):
    clases = Clase.objects.all().order_by('grado', 'grupo')
    clase_id = request.GET.get('clase')
    alumnos = Alumno.objects.filter(activo=True).prefetch_related('clases').order_by('apellido', 'nombre')
    if clase_id and clase_id.isdigit():
        alumnos = alumnos.filter(clases__id=clase_id)
    return render(request, 'Monedas/gestionar_monedas.html', {
        'clases': clases,
        'alumnos': alumnos.distinct(),
        'clase_seleccionada': int(clase_id) if clase_id and clase_id.isdigit() else None,
        'seccion': 'economia',
    })


@require_POST
@login_required
def api_actualizar_monedas(request):
    """Recibe peticiones AJAX para sumar o restar monedas."""
    try:
        alumno = get_object_or_404(Alumno, id=request.POST.get('alumno_id'))
        cantidad = abs(int(request.POST.get('cantidad', 0)))
        accion = request.POST.get('accion')
        motivo = request.POST.get('motivo', '').strip() or 'Ajuste del profesor'
        juego.ajustar_monedas(alumno, cantidad if accion == 'sumar' else -cantidad, motivo)
        alumno.refresh_from_db(fields=['monedas'])
        return JsonResponse({
            'status': 'success',
            'nuevas_monedas': alumno.monedas,
            'mensaje': f'Monedas actualizadas para {alumno.nombre}'
        })
    except (ValueError, TypeError) as error:
        return JsonResponse({'status': 'error', 'mensaje': str(error)}, status=400)


@require_POST
@login_required
def api_monedas_masivo(request):
    try:
        cantidad = abs(int(request.POST.get('cantidad', 0)))
    except ValueError:
        return JsonResponse({'status': 'error', 'mensaje': 'Cantidad inválida'}, status=400)
    accion = request.POST.get('accion')
    motivo = request.POST.get('motivo', '').strip() or 'Premio grupal'
    alumnos = Alumno.objects.filter(activo=True)
    if request.POST.get('clase', '').isdigit():
        alumnos = alumnos.filter(clases__id=request.POST['clase'])
    resultado = {}
    for alumno in alumnos.distinct():
        juego.ajustar_monedas(alumno, cantidad if accion == 'sumar' else -cantidad, motivo)
        resultado[alumno.id] = alumno.monedas
    return JsonResponse({'status': 'success', 'alumnos': resultado, 'total': len(resultado)})


@login_required
def actividad(request):
    movimientos = Movimiento.objects.select_related('alumno').order_by('-fecha')
    tipo = request.GET.get('tipo')
    if tipo:
        movimientos = movimientos.filter(tipo=tipo)
    if request.GET.get('clase', '').isdigit():
        movimientos = movimientos.filter(alumno__clases__id=request.GET['clase']).distinct()
    return render(request, 'Profesor/actividad.html', {
        'pagina': Paginator(movimientos, 50).get_page(request.GET.get('page')),
        'tipos': Movimiento.TIPOS,
        'tipo': tipo,
        'clases': Clase.objects.order_by('grado', 'grupo'),
        'seccion': 'actividad',
    })


# ========================
# MISIONES DIARIAS Y LOGROS
# ========================
@login_required
def diarias_lista(request):
    hoy = timezone.localdate()
    diarias = MisionDiaria.objects.select_related('clase').annotate(
        reclamadas_hoy=Count('progresos', filter=Q(progresos__fecha=hoy, progresos__reclamada=True))
    )
    hoy_por_clase = [
        {'clase': clase, 'diarias': juego.seleccionar_diarias(clase, hoy)}
        for clase in Clase.objects.filter(activa=True).order_by('grado', 'grupo')
    ]
    return render(request, 'Profesor/diarias.html', {
        'diarias': diarias,
        'hoy_por_clase': hoy_por_clase,
        'por_dia': juego.DIARIAS_POR_DIA,
        'seccion': 'diarias',
    })


@login_required
def diaria_form(request, pk=None):
    diaria = get_object_or_404(MisionDiaria, pk=pk) if pk else None
    inicial = {} if diaria else {'objetivo': 'misiones', 'meta': 2, 'recompensa_monedas': 15, 'recompensa_xp': 15,
                                 'icono': '🗺️', 'activa': True}
    form = MisionDiariaForm(request.POST or None, instance=diaria, initial=inicial)
    if request.method == 'POST' and form.is_valid():
        diaria = form.save()
        return _guardado(request, f'Misión del día "{diaria.titulo}" guardada.', reverse('diarias_lista'))
    return _formulario(request, 'Profesor/diaria_form.html', {
        'form': form,
        'diaria': diaria,
        'objetivos': OBJETIVOS_DIARIA,
        'iconos': ICONOS_DIARIA,
        'por_dia': juego.DIARIAS_POR_DIA,
        'activas': MisionDiaria.objects.filter(activa=True).count(),
        'seccion': 'diarias',
    })


@login_required
def logros_lista(request):
    logros = Logro.objects.annotate(total=Count('alumnos'))
    return render(request, 'Profesor/logros.html', {
        'logros': logros,
        'form_otorgar': OtorgarLogroForm(),
        'recientes': LogroAlumno.objects.select_related('alumno', 'logro')[:10],
        'seccion': 'logros',
    })


@login_required
def logro_form(request, pk=None):
    logro = get_object_or_404(Logro, pk=pk) if pk else None
    inicial = {} if logro else {'condicion': 'misiones', 'valor': 5, 'rareza': 'comun', 'icono': '🏅',
                                 'recompensa_monedas': 10, 'recompensa_xp': 0, 'activo': True}
    form = LogroForm(request.POST or None, instance=logro, initial=inicial)
    if request.method == 'POST' and form.is_valid():
        logro = form.save()
        entregados = juego.entregar_logro_a_quienes_cumplen(logro)
        mensaje = f'Logro "{logro.nombre}" guardado.'
        if entregados:
            mensaje += (' 1 alumno ya lo desbloqueó al instante. 🎉' if entregados == 1
                        else f' {entregados} alumnos ya lo desbloquearon al instante. 🎉')
        return _guardado(request, mensaje, reverse('logros_lista'))
    return _formulario(request, 'Profesor/logro_form.html', {
        'form': form,
        'logro': logro,
        'condiciones': CONDICIONES_INFO,
        'iconos': ICONOS_LOGRO,
        'total_alumnos': Alumno.objects.filter(activo=True).count(),
        'seccion': 'logros',
    })


@login_required
def api_logro_cumplen(request):
    condicion = request.GET.get('condicion', '')
    return JsonResponse({'total': len(juego.alumnos_que_cumplen(condicion, request.GET.get('valor')))})


# ========================
# MODO PROYECTOR
# ========================
@login_required
def proyector(request, clase_id):
    clase = get_object_or_404(Clase, id=clase_id)
    semana = juego.ranking_clase(clase, 'semana')
    total = juego.ranking_clase(clase, 'total')
    alumnos = [a for a in total]
    mundos = []
    for mundo in clase.mundos.all():
        ids = list(Mision.objects.filter(nivel__mundo=mundo, activa=True).values_list('id', flat=True))
        posibles = len(ids) * len(alumnos)
        hechas = ProgresoMision.objects.filter(mision_id__in=ids, estado='aprobado', alumno__in=alumnos).count()
        mundos.append({'mundo': mundo, 'tema': juego.TEMA_INFO.get(mundo.tema),
                       'pct': int(hechas * 100 / posibles) if posibles else 0})
    return render(request, 'Profesor/proyector.html', {
        'clase': clase,
        'podio': semana[:3],
        'semana': semana[3:10],
        'total': total[:10],
        'mundos': mundos,
        'logros': LogroAlumno.objects.filter(alumno__clases=clase).select_related('alumno', 'logro')
        .order_by('-fecha').distinct()[:6],
        'equipos': juego.ranking_equipos(clase),
        'metas': juego.metas_visibles(clase),
    })


@login_required
def cargar_niveles(request):
    clase_id = request.GET.get('clase_id')
    if clase_id:
        niveles = Nivel.objects.filter(clase_id=clase_id).order_by('mundo__orden', 'orden')
        return JsonResponse(list(niveles.values('id', 'nombre')), safe=False)
    return JsonResponse([], safe=False)
